"""Offline tests for the PPT-deck export feature.

Python requirements: the project requirements (python-pptx + Pillow are
enough); no network access is used — the iCourse/WebVPN layer and the mail
transport are stubbed.

    python test_ppt_export.py
"""
import io
import os
import random
import sys
import types

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
os.chdir(ROOT)                      # config paths are relative ("data/...")


# Stubs for network / third-party packages that the tested code imports.
_stub = types.ModuleType("src.api.webvpn")
_stub.WebVPNSession = object
_stub.get_vpn_url = lambda url: url
sys.modules["src.api.webvpn"] = _stub

_md = types.ModuleType("markdown")
_md.markdown = lambda text, **kwargs: f"<p>{text}</p>"
sys.modules["markdown"] = _md

_pygfmt = types.ModuleType("pygments.formatters")


class _HtmlFormatter:
    def __init__(self, *a, **k):
        pass

    def get_style_defs(self, *a, **k):
        return ""


_pygfmt.HtmlFormatter = _HtmlFormatter
sys.modules["pygments"] = types.ModuleType("pygments")
sys.modules["pygments.formatters"] = _pygfmt
sys.modules["requests"] = types.ModuleType("requests")

from PIL import Image, ImageDraw                     # noqa: E402
from pptx import Presentation                        # noqa: E402

from src.api import icourse                          # noqa: E402
from src.pipeline import ppt_exporter                # noqa: E402
from src.pipeline import ppt_slide_cache             # noqa: E402
from src.runtime import config                       # noqa: E402


def _png(width, height, color, text):
    im = Image.new("RGB", (width, height), color)
    ImageDraw.Draw(im).text((40, 40), text, fill=(255, 255, 255))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _noise_png(width, height, seed):
    rnd = random.Random(seed)
    raw = bytes(rnd.getrandbits(8) for _ in range(width * height * 3))
    im = Image.frombytes("RGB", (width, height), raw)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _rows(n):
    return [
        {
            "page_num": i + 1,
            "created_sec": 60 * i,
            "pptimgurl": f"https://cdn.example/slide-{i}.png",
            "text": f"page-{i + 1}-ocr-text",
        }
        for i in range(n)
    ]


_ORIGINAL_DOWNLOAD = icourse.download_ppt_image


def _patch_download(mapping):
    icourse.download_ppt_image = lambda client, item: mapping.get(
        item["pptimgurl"]
    )


class FakeDB:
    def __init__(self, rows):
        self._rows = rows

    def get_ppt_export_pages(self, sub_id):
        return self._rows


def cleanup():
    icourse.download_ppt_image = _ORIGINAL_DOWNLOAD
    ppt_exporter.cleanup_exports()


def test_basic_deck():
    cleanup()
    imgs = {
        f"https://cdn.example/slide-{i}.png": _png(
            1920, 1080, (20 + 20 * i, 60, 120), f"Slide {i}"
        )
        for i in range(5)
    }
    _patch_download(imgs)
    path = ppt_exporter.build_lecture_pptx(
        None, FakeDB(_rows(5)), "sub-1",
        "数字逻辑与部件设计", "2026-10-26第1-3节", "2026-10-26",
    )
    assert path and os.path.exists(path), "deck not produced"
    prs = Presentation(path)
    assert len(prs.slides) == 6, f"expected 6 slides, got {len(prs.slides)}"
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "page-1-ocr-text" in notes, notes
    size_mb = os.path.getsize(path) / 1048576
    assert size_mb < config.PPT_EXPORT_MAX_MB + 1, size_mb
    print(f"PASS basic deck: 5 slides + title, {size_mb:.2f} MB")


def test_size_ladder():
    cleanup()
    imgs = {
        f"https://cdn.example/slide-{i}.png": _noise_png(1600, 900, i)
        for i in range(6)
    }
    _patch_download(imgs)
    old = config.PPT_EXPORT_MAX_MB
    config.PPT_EXPORT_MAX_MB = 0.4
    try:
        path = ppt_exporter.build_lecture_pptx(
            None, FakeDB(_rows(6)), "sub-2", "course", "lecture-1",
            "2026-10-26",
        )
    finally:
        config.PPT_EXPORT_MAX_MB = old
    assert path and os.path.exists(path)
    assert len(Presentation(path).slides) == 7
    size_mb = os.path.getsize(path) / 1048576
    assert size_mb < 1.5, f"size ladder ineffective: {size_mb:.2f} MB"
    print(f"PASS size ladder: 6 noisy slides compressed to {size_mb:.2f} MB "
          f"(budget 0.4 MB)")


def test_empty_and_failures():
    cleanup()
    assert ppt_exporter.build_lecture_pptx(
        None, FakeDB([]), "s", "c", "t", "d") is None
    _patch_download({})
    assert ppt_exporter.build_lecture_pptx(
        None, FakeDB(_rows(3)), "s2", "c", "t", "d") is None
    print("PASS empty / all-download-failed -> None")


def test_robust_download_retry():
    cleanup()

    class Resp:
        def __init__(self, data, expected):
            self._data = data
            self.headers = {"Content-Length": str(expected)}

        def raise_for_status(self):
            pass

        def iter_content(self, chunk_size=65536):
            yield self._data

        def close(self):
            pass

    class VPN:
        def __init__(self, script):
            self.script = list(script)
            self.calls = 0

        def get(self, url, **kwargs):
            self.calls += 1
            return self.script.pop(0)

        def get_raw(self, url, **kwargs):
            self.calls += 1
            return self.script.pop(0)

    class Client:
        pass

    vpn = VPN([Resp(b"short", 120), Resp(b"x" * 120, 120)])
    client = Client()
    client.vpn = vpn
    old_sleep = icourse.time.sleep
    icourse.time.sleep = lambda seconds: None
    try:
        data = icourse.download_ppt_image(
            client, {"pptimgurl": "https://cdn.example/x.png", "page_num": 7})
    finally:
        icourse.time.sleep = old_sleep
    assert data == b"x" * 120, data
    assert vpn.calls == 2, vpn.calls
    print("PASS robust image download: short read retried with backoff")


def test_cache_first():
    cleanup()
    rows = _rows(3)
    slides = {
        row["page_num"]: _png(800, 600, (5 * i, 40, 90), f"cached {i}")
        for i, row in enumerate(rows)
    }
    assert ppt_slide_cache.save_pages("sub-cache", slides) == 3

    def must_not_download(*args, **kwargs):
        raise AssertionError("cached screenshots must be reused")

    old = icourse.download_ppt_image
    icourse.download_ppt_image = must_not_download
    try:
        path = ppt_exporter.build_lecture_pptx(
            None, FakeDB(rows), "sub-cache", "course", "lecture-1",
            "2026-10-26")
    finally:
        icourse.download_ppt_image = old
    assert path and os.path.exists(path)
    assert len(Presentation(path).slides) == 4
    assert not os.path.isdir(ppt_slide_cache.lecture_dir("sub-cache")), \
        "per-lecture slide cache should be dropped after building"
    print("PASS deck built from this run's slide cache (no re-download)")


def test_emailer_attachment():
    import email
    import email.header
    import smtplib

    from src.api import emailer as emailer_mod

    cleanup()
    imgs = {
        f"https://cdn.example/slide-{i}.png": _png(
            1280, 720, (10, 90, 60), f"P{i}")
        for i in range(4)
    }
    _patch_download(imgs)
    deck = ppt_exporter.build_lecture_pptx(
        None, FakeDB(_rows(4)), "sub-3",
        "集合与图论（H）", "2026-10-27第4-5节", "2026-10-27",
    )
    assert deck

    sent = {}

    class FakeSMTP:
        def __init__(self, *a, **k):
            pass

        def login(self, *a, **k):
            pass

        def sendmail(self, sender, receiver, message):
            sent["message"] = message

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    old_smtp = smtplib.SMTP_SSL
    smtplib.SMTP_SSL = FakeSMTP
    config.SMTP_EMAIL = "sender@example.com"
    config.SMTP_PASSWORD = "pw"
    config.RECEIVER_EMAIL = "receiver@example.com"
    try:
        ok = emailer_mod.Emailer().send([{
            "sub_id": "sub-3",
            "course_title": "集合与图论（H）",
            "sub_title": "2026-10-27第4-5节",
            "date": "2026-10-27",
            "summary": "hello slides",
            "pptx_path": deck,
        }])
    finally:
        smtplib.SMTP_SSL = old_smtp
        cleanup()

    assert ok, "emailer reported failure"
    parsed = email.message_from_string(sent["message"])
    subject = str(email.header.make_header(
        email.header.decode_header(parsed["Subject"])))
    assert "PPT 课件" in subject, subject
    filenames = [
        part.get_filename()
        for part in parsed.walk()
        if part.get_filename()
    ]
    assert any(f.endswith("_slides.pptx") for f in filenames), filenames
    assert any(f.endswith("_summaries.md") for f in filenames), filenames
    print(f"PASS emailer attachments: {filenames}")


def test_emailer_budget_skip():
    import email
    import email.header
    import smtplib

    from src.api import emailer as emailer_mod

    cleanup()
    os.makedirs(config.PPT_EXPORT_DIR, exist_ok=True)
    deck = os.path.join(config.PPT_EXPORT_DIR, "big.pptx")
    with open(deck, "wb") as f:
        f.write(b"0" * 2048)

    old_budget = config.EMAIL_MAX_ATTACHMENT_MB
    config.EMAIL_MAX_ATTACHMENT_MB = 0.0001
    old_smtp = smtplib.SMTP_SSL
    sent = {}

    class FakeSMTP:
        def __init__(self, *a, **k):
            pass

        def login(self, *a, **k):
            pass

        def sendmail(self, sender, receiver, message):
            sent["message"] = message

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    smtplib.SMTP_SSL = FakeSMTP
    try:
        ok = emailer_mod.Emailer().send([{
            "sub_id": "s", "course_title": "c", "sub_title": "t",
            "date": "d", "summary": "x", "pptx_path": deck,
        }])
    finally:
        smtplib.SMTP_SSL = old_smtp
        config.EMAIL_MAX_ATTACHMENT_MB = old_budget
        cleanup()
    assert ok
    parsed = email.message_from_string(sent["message"])
    subject = str(email.header.make_header(
        email.header.decode_header(parsed["Subject"])))
    assert "PPT 课件" not in subject, subject
    names = [p.get_filename() for p in parsed.walk() if p.get_filename()]
    assert not any((n or "").endswith(".pptx") for n in names), names
    print("PASS emailer budget guard (deck dropped consistently)")


if __name__ == "__main__":
    test_basic_deck()
    test_size_ladder()
    test_empty_and_failures()
    test_robust_download_retry()
    test_cache_first()
    test_emailer_attachment()
    test_emailer_budget_skip()
    print("ALL PPT TESTS PASSED")
